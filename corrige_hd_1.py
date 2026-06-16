"""
Correção HD-1: Alinha PIU_MOBILE_HD.xlsx à tabela mãe.
Gera PIU_MOBILE_HD_CORRIGIDO.xlsx
Alteração cirúrgica: apenas coluna 6. B/D preservados.
"""
import openpyxl, shutil, re

SRC = "PIU_MOBILE_HD.xlsx"
DST = "PIU_MOBILE_HD_CORRIGIDO.xlsx"
shutil.copy2(SRC, DST)
wb = openpyxl.load_workbook(DST, data_only=False)
ws = wb["Itens - Regras de Preço"]

U = '/root/.claude/uploads/311cc8c5-e636-5ebe-9ac9-ac0347d05c63/'

# Carregar preços CAATI e REGIA do PIURB
piurb_wb = openpyxl.load_workbook(U+'8042adc0-TABELA_DE_PRE_O___PIURB_2026__HD.xlsx', data_only=True)
def piurb_row(target):
    for sh in piurb_wb.sheetnames:
        cur = {}
        for row in piurb_wb[sh].iter_rows(values_only=True):
            n_ = row[0] if row else None
            if isinstance(n_, str):
                up_ = n_.strip().upper()
                if up_ == 'PRODUTO':
                    cur = {ci: v.strip() for ci, v in enumerate(row) if isinstance(v, str) and v.strip()}
                    continue
                if up_ == target.upper():
                    return {cur.get(ci, f'c{ci}'): row[ci] for ci in range(len(row)) if isinstance(row[ci], (int, float))}
    return {}

GMAP = {'FX FORN. RB':'FORN','TC FORN. RB':'FORN','FX G':'G','FX H':'H',
        'FX I':'I','FX J':'J/N','FX J / N':'J/N','FX R':'R','FX V':'V',
        'CR VQ':'CR VQ','CR FZ':'CR FZ','COUR FORN.':'COURO FORN'}

def piurb_prices(target):
    raw = piurb_row(target)
    return {GMAP[k]: v for k, v in raw.items() if k in GMAP}

caati_pol = piurb_prices('CAATI POL 0,87X0,95')
caati_puf = piurb_prices('CAATI PUFF 0,65X0,40')
regia_pol = piurb_prices('REGIA POL 0,70X0,73')
print("CAATI POL:", caati_pol)
print("CAATI PUF:", caati_puf)
print("REGIA POL:", regia_pol)

n_fix = 0
def setv(r, v, tag):
    global n_fix
    cur = ws.cell(r, 6).value
    if cur != v:
        ws.cell(r, 6).value = v; n_fix += 1
        print(f"  {tag}: {cur} -> {v}")

def get_grade(cond):
    """Extrai o valor da grade (tecido) da condição."""
    m = re.search(r'(?:TECIDO_\w+|CAT_TECIDO\w*)="([^"]+)"', cond)
    return m.group(1) if m else None

def get_dim(cond):
    m = re.search(r'DIMENSAO_\w+="([^"]+)"', cond)
    return m.group(1) if m else None

# Mapeamento grade Focco -> chave piurb_prices para CAATI
CAATI_GRADE = {
    'FX G': 'G', 'FX H': 'H', 'FX I': 'I', 'FX J/N': 'J/N',
    'FX R': 'R', 'FX V': 'V', 'CR VQ': 'CR VQ', 'CR FZ': 'CR FZ',
    'TECIDO_N': 'J/N',  # TECIDO_N mapeia para FX J/N
}
REGIA_GRADE = {
    'FORNECIDO': 'FORN', 'TECIDO_N': 'J/N', 'TECIDO_R': 'R', 'TECIDO_V': 'V',
}

for r in range(2, ws.max_row + 1):
    cod = str(ws.cell(r, 1).value or '')
    cond = str(ws.cell(r, 3).value or '')
    grade = get_grade(cond)
    dim = get_dim(cond)

    # ── FORNECIDO bug de metragem ──────────────────────────────
    if grade in ('FX FORNECIDO', 'FORNECIDO') and ws.cell(r, 6).value is not None and ws.cell(r, 6).value < 500:
        val = ws.cell(r, 6).value
        # só corrige se é claramente metragem (valor < 200)
        if val < 200:
            fornmap = {
                ('17071','190X100X090'): 4947,
                ('17071','240X100X090'): 5525,
                ('17071','290X100X090'): 6270,
                ('17073',None): 4146,
                ('17074',None): 2255,
                ('17075','200X103X087'): 5113,
                ('17075','220X103X087'): 5319,
                ('17075','240X103X087'): 5615,
                ('17075','260X103X087'): 5941,
                ('17075','280X103X087'): 6227,
                ('17075','300X103X087'): 6480,
                ('17079',None): 3722,
                ('17092','210X095X085'): 5073,
                ('17092','230X095X085'): 5294,
                ('17092','250X095X085'): 5516,
                ('17092','270X095X085'): 5699,
                ('17092','290X095X085'): 5952,
                ('17093','200X102X085'): 4703.35,
                ('17093','220X102X085'): 4919.17,
                ('17093','240X102X085'): 5136.08,
                ('17093','260X102X085'): 5478.34,
                ('17093','280X102X085'): 5731.22,
                ('17093','300X102X085'): 5948.13,
                ('17094',None): 4485,
                ('17095','150X100X090'): 4715,
                ('17095','200X100X090'): 5284,
                ('17095','250X100X090'): 5962,
                ('17096','150X133X090'): 5016,
                ('17096','200X133X090'): 5663,
                ('17096','250X133X090'): 6164,
            }
            key = (cod, dim) if (cod, dim) in fornmap else (cod, None)
            if key in fornmap:
                setv(r, fornmap[key], f"{cod} {dim} FORN-BUG")
                continue

    # ── MESC AP/AV FORNECIDO ──────────────────────────────────
    if cod == '17072' and grade in ('FX FORNECIDO','FORNECIDO'):
        setv(r, 5591, "17072 MESC AV FORN")
    elif cod == '17076' and grade in ('FX FORNECIDO','FORNECIDO'):
        setv(r, 5559, "17076 MESC AP FORN")

    # ── CAATI POL 17011 ───────────────────────────────────────
    elif cod == '17011' and grade in CAATI_GRADE:
        k = CAATI_GRADE[grade]
        if k in caati_pol:
            setv(r, caati_pol[k], f"17011 CAATI {grade}")

    # ── CAATI PUF 17012 ───────────────────────────────────────
    elif cod == '17012' and grade in CAATI_GRADE:
        k = CAATI_GRADE[grade]
        if k in caati_puf:
            setv(r, caati_puf[k], f"17012 CAATI PUF {grade}")

    # ── REGIA POL 17006 ───────────────────────────────────────
    elif cod == '17006' and grade in REGIA_GRADE:
        k = REGIA_GRADE[grade]
        if k in regia_pol:
            setv(r, regia_pol[k], f"17006 REGIA {grade}")

    # ── CACAU POL 17025 (Italo) ───────────────────────────────
    elif cod == '17025':
        if grade == 'TEC_FORNECIDO': setv(r, 3804, "17025 CACAU FORN")
        elif grade == 'TECIDO_N':    setv(r, 4490, "17025 CACAU N")

    # ── CACAU PUF 17030 (Italo) ───────────────────────────────
    elif cod == '17030':
        if grade == 'TECIDO_FORNECIDO': setv(r, 2099, "17030 CACAU PUF FORN")
        elif grade == 'TECIDO_N':        setv(r, 2351, "17030 CACAU PUF N")

    # ── PIURB FORNECIDO ───────────────────────────────────────
    elif cod == '17049' and grade == 'FORNECIDO':
        setv(r, 4165, "17049 CAJU FORN")
    elif cod == '17052' and grade == 'FORNECIDO':
        setv(r, 7862, "17052 CRIO FORN")
    elif cod == '17057' and grade == 'FORNECIDO':
        setv(r, 3961, "17057 TAPIOCA FORN")

    # ── XIQUE XIQUE 17070 FORNECIDO ───────────────────────────
    elif cod == '17070' and grade == 'FX FORNECIDO':
        setv(r, 4128, "17070 XIQUE FORN")

    # ── TULIPA 17067 FX G ─────────────────────────────────────
    elif cod == '17067' and grade == 'G':
        setv(r, 2709, "17067 TULIPA G")

    # ── SAPUCAIA 17081 FX G ───────────────────────────────────
    elif cod == '17081' and grade == 'FX G':
        setv(r, 4567, "17081 SAP FX G")

    # ── MARAA FORN RB ─────────────────────────────────────────
    elif cod == '17084' and grade == 'FX FORN RB':
        setv(r, 5232, "17084 MARAA CTO FORN RB")
    elif cod == '17085' and grade == 'FX FORN RB':
        setv(r, 5204, "17085 MARAA PUFF FORN RB")

    # ── KIM 17198 FORNECIDO ───────────────────────────────────
    elif cod == '17198' and grade == 'FORNECIDO':
        kim_forn = {'0,93X2,08':4662,'1,43X2,08':5097,'1,63X2,18':5283,'1,83X2,18':5445,'1,98X2,23':5571}
        if dim:
            for d, v in kim_forn.items():
                if d in dim: setv(r, v, f"17198 KIM {d} FORN"); break

    # ── SAPUCAIA 17181 FORNECIDO ──────────────────────────────
    elif cod == '17181' and grade == 'FORNECIDO':
        sap_forn = {'0,98X2,22':4021,'1,48X2,22':4429,'1,68X2,32':4657,'1,88X2,32':4879,'2,03X2,37':5092}
        if dim:
            for d, v in sap_forn.items():
                if d in dim: setv(r, v, f"17181 SAP {d} FORN"); break

print(f"\nTotal alterações: {n_fix}")

def _ult(w, ncols):
    u = w.max_row
    while u > 1 and all(w.cell(u, c).value is None for c in range(1, ncols+1)):
        u -= 1
    return u

for sheet, tab, nc, L in [
    ("Lista de Opções","Tabela_ListaOpcoes",5,"E"),
    ("Características","Tabela_Caracteristicas",4,"D"),
    ("Itens","Tabela_Itens",14,"N"),
    ("Itens - Atributos","Tabela_ItensAtributos",3,"C"),
    ("Itens - Regras de Preço","Tabela_ItensRegrasPreco",7,"G")]:
    w = wb[sheet]; w.tables[tab].ref = f"A1:{L}{_ult(w, nc)}"

wb.save(DST)
print(f"✓ Saved {DST}")
