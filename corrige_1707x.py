"""
Correção 4: completa 17071 (POLTRONA CARCARA), 17072 (POLTRONA MABO),
            17078 (PUFF SINHO)
- Todos usam TECIDO_17070 compartilhado (padrão da família HEBE)
- 17071: só tinha FORNECIDO (correto=3722), faltavam 13 faixas
- 17072: tinha FX J/N errado (3480→4249), faltavam 13 faixas
- 17078: tinha FX H correto (2439), faltavam 13 faixas

Tabela mãe (PIU - DESIGNERS):
  CARCARA POL 0,94X0,88 -> FORN=3722 FXA=3798 ... FXKNIT=4831
  MABO POL 0,80X0,88    -> FORN=3774 FXA=3849 ... FXKNIT=4874
  SINHO PUFF 0,80X0,82  -> FORN=2255 FXA=2301 ... FXKNIT=2929

Trabalha sobre PIU_MOBILE_HDA_CORRIGIDO3.xlsx -> PIU_MOBILE_HDA_CORRIGIDO4.xlsx
"""

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color
import shutil

SRC = "PIU_MOBILE_HDA_CORRIGIDO3.xlsx"
DST = "PIU_MOBILE_HDA_CORRIGIDO4.xlsx"
shutil.copy2(SRC, DST)

wb  = openpyxl.load_workbook(DST, data_only=False)
wsRP = wb["Itens - Regras de Preço"]

FILL = PatternFill(patternType="solid", fgColor=Color(theme=0, tint=-0.249977111117893))

def _ult(ws, ncols):
    u = ws.max_row
    while u > 1 and all(ws.cell(u, c).value is None for c in range(1, ncols + 1)):
        u -= 1
    return u

def set_regra(ws, r, cod, cond, valor):
    ws.cell(r, 1, cod)
    b = ws.cell(r, 2, "QUANDO"); b.fill = FILL
    ws.cell(r, 3, cond)
    d = ws.cell(r, 4, "ENTÃO"); d.fill = FILL
    ws.cell(r, 5, "CONCEDE ACRÉSCIMO DE")
    ws.cell(r, 6, valor)
    ws.cell(r, 7, "NO PREÇO BASE")

# Preços da tabela mãe (PIU - DESIGNERS)
precos = {
    '17071': {   # CARCARA POL 0,94X0,88
        'FORNECIDO': 3722, 'FX A': 3798, 'FX B': 3823, 'FX C': 3848,
        'FX D': 3884, 'FX E': 3909, 'FX F': 3949, 'FX G': 3984,
        'FX H': 4025, 'FX I': 4075, 'FX J/N': 4201, 'FX R': 4251,
        'FX V': 4352, 'FX KNIT': 4831,
    },
    '17072': {   # MABO POL 0,80X0,88
        'FORNECIDO': 3774, 'FX A': 3849, 'FX B': 3874, 'FX C': 3899,
        'FX D': 3934, 'FX E': 3959, 'FX F': 3999, 'FX G': 4034,
        'FX H': 4074, 'FX I': 4124, 'FX J/N': 4249, 'FX R': 4299,
        'FX V': 4399, 'FX KNIT': 4874,
    },
    '17078': {   # SINHO PUFF 0,80X0,82
        'FORNECIDO': 2255, 'FX A': 2301, 'FX B': 2317, 'FX C': 2332,
        'FX D': 2353, 'FX E': 2369, 'FX F': 2393, 'FX G': 2414,
        'FX H': 2439, 'FX I': 2470, 'FX J/N': 2546, 'FX R': 2577,
        'FX V': 2638, 'FX KNIT': 2929,
    },
}

FAIXAS_ORDER = ['FORNECIDO','FX A','FX B','FX C','FX D','FX E','FX F',
                'FX G','FX H','FX I','FX J/N','FX R','FX V','FX KNIT']

for cod, p in precos.items():
    # Find and fix/overwrite the existing (single) rule
    linha_existente = None
    existing_faixa = None
    for r in range(2, wsRP.max_row + 1):
        if str(wsRP.cell(r, 1).value or '') == cod:
            cond_old = str(wsRP.cell(r, 3).value or '')
            # detect which faixa it has
            for fx in FAIXAS_ORDER:
                if fx in cond_old or fx.replace('FORNECIDO','FORNECIDO') in cond_old:
                    existing_faixa = fx
                    break
            linha_existente = r
            break

    if linha_existente is None:
        print(f"  ERRO: {cod} sem regra existente!")
        continue

    # Rewrite existing rule as FORNECIDO (first faixa)
    cond_forn = f'TECIDO_17070="FORNECIDO"'
    set_regra(wsRP, linha_existente, cod, cond_forn, p['FORNECIDO'])
    print(f"  {cod}: regra existente ({existing_faixa}) -> FORNECIDO={p['FORNECIDO']}")

    # Append remaining faixas
    r_next = _ult(wsRP, 7) + 1
    for faixa in FAIXAS_ORDER[1:]:
        cond = f'TECIDO_17070="{faixa}"'
        set_regra(wsRP, r_next, cod, cond, p[faixa])
        print(f"  {cod}: + {faixa}={p[faixa]}")
        r_next += 1

# Update table refs
TABELAS = [
    ("Lista de Opções",          "Tabela_ListaOpcoes",      5, "E"),
    ("Características",          "Tabela_Caracteristicas",  4, "D"),
    ("Itens",                    "Tabela_Itens",           14, "N"),
    ("Itens - Atributos",        "Tabela_ItensAtributos",   3, "C"),
    ("Itens - Regras de Preço",  "Tabela_ItensRegrasPreco", 7, "G"),
]
for sheet, tab, nc, L in TABELAS:
    ws = wb[sheet]
    new_ref = f"A1:{L}{_ult(ws, nc)}"
    ws.tables[tab].ref = new_ref

wb.save(DST)
print(f"\n✓ Saved {DST}")
