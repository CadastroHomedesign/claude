"""
Correção 5: divergências confirmadas no audit v2 vs tabelas mãe.
Trabalha sobre PIU_MOBILE_HDA_CORRIGIDO4.xlsx -> PIU_MOBILE_HDA_CORRIGIDO5.xlsx

Grupos:
1. 22113 CHAISE MESC CSB (= MESC CSB 1,50X1,66): todas as faixas defasadas.
   - Linha J/N com dimensão (6656) já correta -> mantém.
   - Linha ESPECIAL (130X166) -> ignora.
2. 22114 / 22115 (AP/AV): só FORNECIDO errado (bug metragem). Faixas têm
   offset constante (+171 / +203) -> FORNECIDO = 5388 + offset.
3. PIURB FORNECIDO desatualizado: 22011, 22012, 22032, 22033, 22041.
4. 22028 SOFA CAMA PALERMO 2,30 (dim principal): todas as faixas.
   - Dim ESPECIAL -> ignora.
"""

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color
import shutil

SRC = "PIU_MOBILE_HDA_CORRIGIDO4.xlsx"
DST = "PIU_MOBILE_HDA_CORRIGIDO5.xlsx"
shutil.copy2(SRC, DST)

wb = openpyxl.load_workbook(DST, data_only=False)
ws = wb["Itens - Regras de Preço"]
FILL = PatternFill(patternType="solid", fgColor=Color(theme=0, tint=-0.249977111117893))

def fix_qe(r):
    # Preserva B/D originais (formato/fill do arquivo). Nada a fazer:
    # estas correções alteram apenas a coluna 6 (preço).
    return

def _ult(w, ncols):
    u = w.max_row
    while u > 1 and all(w.cell(u, c).value is None for c in range(1, ncols + 1)):
        u -= 1
    return u

# ── Grupo 1: 22113 (chave = faixa em TECIDO_22113), valores = MESC CSB 1,50X1,66
p22113 = {
    'FX FORNECIDO': 5388, 'FX A': 5588, 'FX B': 5655, 'FX C': 5722,
    'FX D': 5815, 'FX E': 5882, 'FX F': 5989, 'FX G': 6082,
    'FX H': 6189, 'FX I': 6322, 'FX J/N': 6656, 'FX R': 6789,
    'FX V': 7056, 'FX KNIT': 8323,
}

# ── Grupo 3: PIURB FORNECIDO  (cod -> novo valor)
forn_piurb = {'22011': 3847, '22012': 2286, '22032': 6606, '22033': 4165, '22041': 3961}

# ── Grupo 4: 22028 PALERMO dim principal (TECIDO_x -> valor)
p22028 = {
    'TECIDO_A': 4756, 'TECIDO_B': 4847, 'TECIDO_C': 4937, 'TECIDO_D': 5064,
    'TECIDO_E': 5154, 'TECIDO_F': 5299, 'TECIDO_G': 5425, 'TECIDO_H': 5570,
    'TECIDO_I': 5751, 'TECIDO_J': 6202,
}

# ── Grupo 2: FORNECIDO de 22114/22115 com offset constante sobre base 5388
forn_apav = {'22114': 5388 + 171, '22115': 5388 + 203}  # 5559 / 5591

n = 0
for r in range(2, ws.max_row + 1):
    cod = str(ws.cell(r, 1).value or '')
    cond = str(ws.cell(r, 3).value or '')

    # Grupo 1: 22113 sem dimensão na condição (ignora J/N com dim e ESPECIAL)
    if cod == '22113' and 'DIMENSAO' not in cond:
        for fx, v in p22113.items():
            if f'TECIDO_22113="{fx}"' in cond:
                if ws.cell(r, 6).value != v:
                    ws.cell(r, 6).value = v; fix_qe(r); n += 1
                    print(f"  22113 {fx} -> {v}")
                break
        continue

    # Grupo 3: PIURB FORNECIDO
    if cod in forn_piurb and 'FORNECIDO' in cond.upper():
        v = forn_piurb[cod]
        if ws.cell(r, 6).value != v:
            ws.cell(r, 6).value = v; fix_qe(r); n += 1
            print(f"  {cod} FORNECIDO -> {v}")
        continue

    # Grupo 4: 22028 dim principal
    if cod == '22028' and '2,30X0,95/1,32X0,80"' in cond:
        for tk, v in p22028.items():
            if f'CAT_TECIDO03="{tk}"' in cond:
                if ws.cell(r, 6).value != v:
                    ws.cell(r, 6).value = v; fix_qe(r); n += 1
                    print(f"  22028 {tk} -> {v}")
                break
        continue

    # Grupo 2: 22114/22115 FORNECIDO
    if cod in forn_apav and 'FORNECIDO' in cond.upper():
        v = forn_apav[cod]
        if ws.cell(r, 6).value != v:
            ws.cell(r, 6).value = v; fix_qe(r); n += 1
            print(f"  {cod} FX FORNECIDO -> {v}")
        continue

print(f"\nTotal células corrigidas: {n}")

# Refs (nenhuma linha inserida, mas mantém consistência)
TABELAS = [
    ("Lista de Opções", "Tabela_ListaOpcoes", 5, "E"),
    ("Características", "Tabela_Caracteristicas", 4, "D"),
    ("Itens", "Tabela_Itens", 14, "N"),
    ("Itens - Atributos", "Tabela_ItensAtributos", 3, "C"),
    ("Itens - Regras de Preço", "Tabela_ItensRegrasPreco", 7, "G"),
]
for sheet, tab, nc, L in TABELAS:
    w = wb[sheet]
    w.tables[tab].ref = f"A1:{L}{_ult(w, nc)}"

wb.save(DST)
print(f"✓ Saved {DST}")
