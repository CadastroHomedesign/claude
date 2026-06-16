import openpyxl
wb=openpyxl.load_workbook('PIU_MOBILE_HD_CADASTRO.xlsx',data_only=False)
R=wb['Itens - Regras de Preço']
def info(code):
    for r in range(2,R.max_row+1):
        if str(R.cell(r,1).value)==code:
            b=R.cell(r,2); d=R.cell(r,4)
            patt = b.fill.patternType if b.fill else None
            theme = b.fill.fgColor.theme if (b.fill and b.fill.patternType) else None
            tint = b.fill.fgColor.tint if (b.fill and b.fill.patternType) else None
            return (repr(b.value), repr(d.value), patt, theme, tint)
    return None
print('22370 novo     :', info('22370'))
print('17072 existente:', info('17072'))
