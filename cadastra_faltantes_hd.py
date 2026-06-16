"""
Cadastra no HD os 24 produtos que existem no HDA e faltavam no HD.
- 12 produtos cujo codigo 17xxx ja existe no HD (colisao) -> novos codigos 17469..
- 12 produtos 22xxx (livres no HD) -> mantem o codigo
Cada produto fica AUTOCONTIDO: toda caracteristica ESCOLHA recebe nome unico
(<nome>__<codigo_final>) e a lista de opcoes e copiada do HDA. Assim nao colide
com caracteristicas homonimas ja existentes no HD.
Precos vem do HDA (ja alinhado as tabelas mae). Preco Base = 0. B/D preservados.
Gera PIU_MOBILE_HD_CADASTRO.xlsx
"""
import openpyxl, re, copy
from openpyxl.styles import PatternFill

GRAY = PatternFill(patternType="solid", fgColor="FFBFBFBF")

SRC = "PIU_MOBILE_HD_CORRIGIDO.xlsx"
DST = "PIU_MOBILE_HD_CADASTRO.xlsx"
import shutil; shutil.copy2(SRC, DST)

hd  = openpyxl.load_workbook(DST,  data_only=False)
hda = openpyxl.load_workbook("PIU_MOBILE_HDA_CORRIGIDO7.xlsx", data_only=False)

# ---- mapeamento de codigos ----
collide = ['17066','17069','17073','17074','17075','17079',
           '17082','17084','17085','17086','17087','17088']
keep    = ['22018','22027','22036','22044',
           '22370','22371','22372','22373','22374','22375','22376','22377']
newcode = {oc: str(17469+i) for i, oc in enumerate(collide)}
for k in keep: newcode[k] = k          # mantem
ALL = collide + keep

# ---- indices das abas HDA ----
H_ITENS = hda['Itens']; H_ATR = hda['Itens - Atributos']
H_REG   = hda['Itens - Regras de Preço']; H_CAR = hda['Características']
H_OPC   = hda['Lista de Opções']

# tipo e lista de cada caracteristica no HDA
hda_char = {}
for r in range(2, H_CAR.max_row+1):
    n = H_CAR.cell(r,1).value
    if n: hda_char[str(n)] = (H_CAR.cell(r,2).value, H_CAR.cell(r,3).value, H_CAR.cell(r,4).value)

# opcoes por lista (nome da lista -> [ (id,desc,ativa,tags) ])
from collections import defaultdict
hda_opts = defaultdict(list)
for r in range(2, H_OPC.max_row+1):
    d = H_OPC.cell(r,1).value
    if d: hda_opts[str(d)].append([H_OPC.cell(r,c).value for c in range(2,6)])

# linhas HDA por codigo
def rows_for(ws, code, ncol):
    return [[ws.cell(r,c).value for c in range(1,ncol+1)] for r in range(2,ws.max_row+1)
            if str(ws.cell(r,1).value)==code]
def rowobjs_for(ws, code):
    return [r for r in range(2,ws.max_row+1) if str(ws.cell(r,1).value)==code]

# caracteristicas ja existentes no HD
hd_char_names = set()
HC = hd['Características']
for r in range(2, HC.max_row+1):
    n = HC.cell(r,1).value
    if n: hd_char_names.add(str(n))

# abas destino
D_ITENS = hd['Itens']; D_ATR = hd['Itens - Atributos']
D_REG = hd['Itens - Regras de Preço']; D_CAR = hd['Características']; D_OPC = hd['Lista de Opções']

def _ult(w, nc):
    u=w.max_row
    while u>1 and all(w.cell(u,c).value is None for c in range(1,nc+1)): u-=1
    return u

# ponteiros de escrita
p_it  = _ult(D_ITENS,14)+1
p_atr = _ult(D_ATR,3)+1
p_reg = _ult(D_REG,7)+1
p_car = _ult(D_CAR,4)+1
p_opc = _ult(D_OPC,5)+1

CHAR_TOK = re.compile(r'([A-Za-z_][A-Za-z0-9_]*)=')
new_text_chars = {}   # nome TEXTO faltante -> desc (criar uma vez)

for oc in ALL:
    nc = newcode[oc]
    # ---- caracteristicas usadas (cols 5..14 do Itens) ----
    it_rowobj = rowobjs_for(H_ITENS, oc)[0]
    char_cols = [H_ITENS.cell(it_rowobj, c).value for c in range(5,15)]
    # tokens tambem usados nas condicoes
    cond_tokens = set()
    for r in rowobjs_for(H_REG, oc):
        cond = str(H_REG.cell(r,3).value or '')
        cond_tokens.update(CHAR_TOK.findall(cond))

    # montar mapa de renome para ESCOLHA; TEXTO mantem (garante existir no HD)
    rename = {}
    for name in list(char_cols)+list(cond_tokens):
        if not name: continue
        name=str(name)
        info = hda_char.get(name)
        if not info: continue
        desc, tipo, lista = info
        if tipo=='ESCOLHA':
            rename[name] = f"{name}__{nc}"
        else:  # TEXTO
            if name not in hd_char_names:
                new_text_chars[name]=desc

    # ---- Itens ----
    base = [H_ITENS.cell(it_rowobj,c).value for c in range(1,15)]
    base[0]=nc; base[2]=0
    for ci in range(4,14):  # cols 5..14 -> idx4..13
        v=base[ci]
        if v in rename: base[ci]=rename[v]
    for c in range(1,15): D_ITENS.cell(p_it,c).value=base[c-1]
    p_it+=1

    # ---- Atributos ----
    for r in rowobjs_for(H_ATR, oc):
        vals=[H_ATR.cell(r,c).value for c in range(1,4)]; vals[0]=nc
        for c in range(1,4): D_ATR.cell(p_atr,c).value=vals[c-1]
        p_atr+=1

    # ---- Regras de Preco (B/D reconstruidos no formato do HD) ----
    for r in rowobjs_for(H_REG, oc):
        cond=str(H_REG.cell(r,3).value or '')
        for old,new in sorted(rename.items(), key=lambda x:-len(x[0])):
            cond=cond.replace(f'{old}=', f'{new}=')
        has_qe = H_REG.cell(r,2).value is not None  # linha tem QUANDO/ENTAO?
        D_REG.cell(p_reg,1).value = nc
        D_REG.cell(p_reg,3).value = cond
        D_REG.cell(p_reg,5).value = H_REG.cell(r,5).value   # E acao
        D_REG.cell(p_reg,6).value = H_REG.cell(r,6).value   # F valor (preco)
        D_REG.cell(p_reg,7).value = H_REG.cell(r,7).value   # G modificador
        if has_qe:
            b=D_REG.cell(p_reg,2); b.value='="QUANDO"'; b.fill=GRAY
            d=D_REG.cell(p_reg,4); d.value='="ENTÃO"';  d.fill=GRAY
        p_reg+=1

    # ---- Caracteristicas ESCOLHA novas + suas opcoes ----
    for old,new in rename.items():
        desc,tipo,lista = hda_char[old]
        D_CAR.cell(p_car,1).value=new; D_CAR.cell(p_car,2).value=desc
        D_CAR.cell(p_car,3).value='ESCOLHA'; D_CAR.cell(p_car,4).value=new
        p_car+=1
        for opt in hda_opts.get(str(lista),[]):
            D_OPC.cell(p_opc,1).value=new
            for j,val in enumerate(opt): D_OPC.cell(p_opc,2+j).value=val
            p_opc+=1

# ---- TEXTO faltantes (uma vez) ----
for name,desc in new_text_chars.items():
    D_CAR.cell(p_car,1).value=name; D_CAR.cell(p_car,2).value=desc
    D_CAR.cell(p_car,3).value='TEXTO'; D_CAR.cell(p_car,4).value=None
    p_car+=1

print(f"Produtos cadastrados: {len(ALL)}")
print(f"TEXTO criados: {list(new_text_chars)}")

# ---- refs das 5 tabelas ----
for sheet,tab,nc_,L in [
    ("Lista de Opções","Tabela_ListaOpcoes",5,"E"),
    ("Características","Tabela_Caracteristicas",4,"D"),
    ("Itens","Tabela_Itens",14,"N"),
    ("Itens - Atributos","Tabela_ItensAtributos",3,"C"),
    ("Itens - Regras de Preço","Tabela_ItensRegrasPreco",7,"G")]:
    w=hd[sheet]; w.tables[tab].ref=f"A1:{L}{_ult(w,nc_)}"

hd.save(DST)
print(f"OK -> {DST}")
print("Mapa de codigos:")
for oc in collide: print(f"  {oc} -> {newcode[oc]}")
