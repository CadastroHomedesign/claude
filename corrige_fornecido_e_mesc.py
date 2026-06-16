"""
Correção 2: Fix FORNECIDO prices e completa 22213 ALOFI
Trabalha sobre PIU_MOBILE_HDA_CORRIGIDO.xlsx e gera PIU_MOBILE_HDA_CORRIGIDO2.xlsx
"""

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color
import shutil

SRC = "PIU_MOBILE_HDA_CORRIGIDO.xlsx"
DST = "PIU_MOBILE_HDA_CORRIGIDO2.xlsx"

shutil.copy2(SRC, DST)

wb = openpyxl.load_workbook(DST, data_only=False)
wsRP = wb["Itens - Regras de Preço"]
wsLO = wb["Lista de Opções"]
wsC  = wb["Características"]
wsI  = wb["Itens"]
wsIA = wb["Itens - Atributos"]

FILL = PatternFill(patternType="solid", fgColor=Color(theme=0, tint=-0.249977111117893))

def fix_cell(ws, r, c, val):
    ws.cell(r, c, val)

def fix_quando_entao(ws, r):
    b = ws.cell(r, 2); b.value = "QUANDO"; b.fill = FILL
    d = ws.cell(r, 4); d.value = "ENTÃO";  d.fill = FILL

def _ult(ws, ncols):
    u = ws.max_row
    while u > 1 and all(ws.cell(u, c).value is None for c in range(1, ncols + 1)):
        u -= 1
    return u

def add_regra(ws, r, cod, cond, valor):
    ws.cell(r, 1, cod)
    b = ws.cell(r, 2, "QUANDO"); b.fill = FILL
    ws.cell(r, 3, cond)
    d = ws.cell(r, 4, "ENTÃO"); d.fill = FILL
    ws.cell(r, 5, "CONCEDE ACRÉSCIMO DE")
    ws.cell(r, 6, valor)
    ws.cell(r, 7, "NO PREÇO BASE")
    return r + 1

def add_opcao_lo(ws, r, caracteristica, valor, descricao="SIM"):
    ws.cell(r, 1, caracteristica)
    ws.cell(r, 2, valor)
    ws.cell(r, 3, valor)
    ws.cell(r, 4, descricao)
    return r + 1

# ────────────────────────────────────────────────────────────
# 1. FIX FORNECIDO PRICES
# ────────────────────────────────────────────────────────────

# Row lookup for each item: maps partial dimension string -> correct FORN price
fornecido_fixes = {
    # 22066 CAMA KIARA PINTURA - col8 (Couro FORN metragem) was used instead of col9 (FORN price)
    '22066': {
        '105X211': 5480,
        '155X211': 5935,
        '175X221': 6119,
        '195X221': 6277,
        '210X226': 6412,
    },
    # 22124 COLISEU CHAISE
    '22124': {
        '192X192': 5978,
    },
    # 22239 CAPADOCIA 2B
    '22239': {
        '200X103': 5113,
        '220X103': 5319,
        '240X103': 5615,
        '260X103': 5941,
        '280X103': 6227,
        '300X103': 6480,
    },
    # 22312 MEPPEL 2B
    '22312': {
        '200X102': 4703.35,
        '220X102': 4919.17,
        '240X102': 5136.08,
        '260X102': 5478.34,
        '280X102': 5731.22,
        '300X102': 5948.13,
    },
    # 22313 MESC SB
    '22313': {
        '150X100': 4715,
        '200X100': 5284,
        '250X100': 5962,
    },
    # 22314 MESC SB AP
    '22314': {
        '150X133': 5016,
        '200X133': 5663,
        '250X133': 6164,
    },
    # 22315 MESC SB AV
    '22315': {
        '190X100': 4947,
        '240X100': 5525,
        '290X100': 6270,
    },
    # 22325 MURRAY 2B 2AS
    '22325': {
        '210X095': 5073,
        '230X095': 5294,
        '250X095': 5516,
        '270X095': 5699,
        '290X095': 5952,
    },
}

count_forn = 0
for r in range(2, wsRP.max_row + 1):
    cod = str(wsRP.cell(r, 1).value or '')
    if cod not in fornecido_fixes:
        continue
    cond = str(wsRP.cell(r, 3).value or '')
    val_cell = wsRP.cell(r, 6)

    # Check if this is a FORNECIDO rule
    is_forn = (
        'FORNECIDO' in cond.upper() or
        'FX FORNECIDO' in cond or
        'TECIDO FORNECIDO' in cond
    )
    if not is_forn:
        continue

    dim_map = fornecido_fixes[cod]
    matched_price = None
    for dim_key, price in dim_map.items():
        if dim_key in cond:
            matched_price = price
            break

    if matched_price is None:
        print(f"  WARNING: {cod} FORN rule not matched: {cond}")
        continue

    val_cell.value = matched_price
    fix_quando_entao(wsRP, r)
    count_forn += 1
    print(f"  Fixed FORN {cod}: {cond[:50]} => {matched_price}")

print(f"\nFixed {count_forn} FORNECIDO rules\n")

# ────────────────────────────────────────────────────────────
# 2. FIX 22066 175X221 - ALL FAIXAS WRONG
# ────────────────────────────────────────────────────────────
# Correct prices from mother table row 83 (KIARA CAMA 1,75X2,21):
kiara_175_prices = {
    'FX FORNECIDO': 6119,  # already fixed above
    'FX A':    6253,
    'FX B':    6298,
    'FX C':    6343,
    'FX D':    6405,
    'FX E':    6450,
    'FX F':    6522,
    'FX G':    6584,
    'FX H':    6656,
    'FX I':    6745,
    'FX J / N': 6969,
    'FX R':    7058,
    'FX V':    7237,
    'FX KNIT': 8087,
}

count_175 = 0
for r in range(2, wsRP.max_row + 1):
    cod = str(wsRP.cell(r, 1).value or '')
    if cod != '22066':
        continue
    cond = str(wsRP.cell(r, 3).value or '')
    if '175' not in cond:
        continue

    # Find which faixa
    matched_faixa = None
    for faixa, price in kiara_175_prices.items():
        if faixa in cond:
            matched_faixa = faixa
            break

    if matched_faixa is None:
        print(f"  WARNING 22066 175: faixa not matched: {cond}")
        continue

    current = wsRP.cell(r, 6).value
    correct = kiara_175_prices[matched_faixa]
    if current != correct:
        wsRP.cell(r, 6).value = correct
        fix_quando_entao(wsRP, r)
        count_175 += 1
        print(f"  Fixed 22066 175X221 {matched_faixa}: {current} => {correct}")

print(f"\nFixed {count_175} rules for 22066 175X221\n")

# ────────────────────────────────────────────────────────────
# 3. COMPLETE 22213 SOFA ALOFI 1BR AV 2,50
# ────────────────────────────────────────────────────────────
# Current: 1 rule TECIDO_22213="TECIDO FORNECIDO" => 33.6 (metragem bug)
# Fix: change value to 7125
# Add: FX A..FX KNIT (13 new options + 13 new rules)

DIM_22213 = 'C2,50XP1,00X0,78A'
COD_22213 = '22213'

# Faixas for ALOFI 1B AV 2,50 (mother table row 38)
# col9=FORN=7125, col10=FXA=7370, col11=FXB=7452, col12=FXC=7534
# col13=FXD=7648, col14=FXE=7730, col15=FXF=7861, col16=FXG=7976
# col17=FXH=8106, col18=FXI=8270, col19=FXJ=8679, col20=FXR=8843
# col21=FXV=9170, col22=FXKNIT=10724
alofi_faixas = [
    ('FX A',   7370),
    ('FX B',   7452),
    ('FX C',   7534),
    ('FX D',   7648),
    ('FX E',   7730),
    ('FX F',   7861),
    ('FX G',   7976),
    ('FX H',   8106),
    ('FX I',   8270),
    ('FX J/N', 8679),
    ('FX R',   8843),
    ('FX V',   9170),
    ('FX KNIT', 10724),
]

# Fix the existing FORNECIDO rule (value 33.6 -> 7125)
forn_fixed = False
for r in range(2, wsRP.max_row + 1):
    cod = str(wsRP.cell(r, 1).value or '')
    if cod != COD_22213:
        continue
    cond = str(wsRP.cell(r, 3).value or '')
    if 'TECIDO FORNECIDO' in cond or 'FORNECIDO' in cond:
        wsRP.cell(r, 6).value = 7125
        fix_quando_entao(wsRP, r)
        forn_fixed = True
        print(f"  Fixed 22213 FORNECIDO: 33.6 => 7125")

# Add new ListaOpcoes for each faixa under TECIDO_22213
# Find last row of Lista de Opções
lo_last = _ult(wsLO, 5)
rLO = lo_last + 1
for (faixa, _) in alofi_faixas:
    wsLO.cell(rLO, 1, 'TECIDO_22213')
    wsLO.cell(rLO, 2, faixa)
    wsLO.cell(rLO, 3, faixa)
    wsLO.cell(rLO, 4, 'SIM')
    print(f"  Added ListaOpcoes TECIDO_22213 / {faixa}")
    rLO += 1

# Add new rules for each faixa
rp_last = _ult(wsRP, 7)
rRP = rp_last + 1
for (faixa, preco) in alofi_faixas:
    cond = f'DIMENSAO_{COD_22213}="{DIM_22213}" E TECIDO_{COD_22213}="{faixa}"'
    rRP = add_regra(wsRP, rRP, COD_22213, cond, preco)
    print(f"  Added 22213 rule: {faixa} => {preco}")

# ────────────────────────────────────────────────────────────
# 4. UPDATE TABLE REFS
# ────────────────────────────────────────────────────────────
TABELAS = [
    ("Lista de Opções",          "Tabela_ListaOpcoes",     5, "E"),
    ("Características",          "Tabela_Caracteristicas", 4, "D"),
    ("Itens",                    "Tabela_Itens",           14, "N"),
    ("Itens - Atributos",        "Tabela_ItensAtributos",  3, "C"),
    ("Itens - Regras de Preço",  "Tabela_ItensRegrasPreco", 7, "G"),
]
for sheet, tab, nc, L in TABELAS:
    ws = wb[sheet]
    new_ref = f"A1:{L}{_ult(ws, nc)}"
    ws.tables[tab].ref = new_ref
    print(f"  Table {tab}: ref -> {new_ref}")

wb.save(DST)
print(f"\n✓ Saved {DST}")
