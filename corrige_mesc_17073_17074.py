"""
Correção 3: completa 17073 (SOFA MESC SB/AV) e 17074 (CHAISE MESC SB)
- Ambos tinham 1 regra com característica ERRADA (TECIDO_17070) e sem DIMENSAO
- Reescreve usando as características próprias (DIMENSAO_17073/TECIDO_17073 etc.)
  e todas as 14 faixas vindas da tabela mãe (PIU - DESIGNERS / MESC)
Trabalha sobre PIU_MOBILE_HDA_CORRIGIDO2.xlsx -> PIU_MOBILE_HDA_CORRIGIDO3.xlsx
"""

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color
import shutil

SRC = "PIU_MOBILE_HDA_CORRIGIDO2.xlsx"
DST = "PIU_MOBILE_HDA_CORRIGIDO3.xlsx"
shutil.copy2(SRC, DST)

wb = openpyxl.load_workbook(DST, data_only=False)
wsRP = wb["Itens - Regras de Preço"]

FILL = PatternFill(patternType="solid", fgColor=Color(theme=0, tint=-0.249977111117893))

def _ult(ws, ncols):
    u = ws.max_row
    while u > 1 and all(ws.cell(u, c).value is None for c in range(1, ncols + 1)):
        u -= 1
    return u

def set_regra(ws, r, cod, cond, valor):
    ws.cell(r, 1, cod)
    b = ws.cell(r, 2, "QUANDO"); b.fill = FILL
    ws.cell(r, 3, cond)
    d = ws.cell(r, 4, "ENTÃO"); d.fill = FILL
    ws.cell(r, 5, "CONCEDE ACRÉSCIMO DE")
    ws.cell(r, 6, valor)
    ws.cell(r, 7, "NO PREÇO BASE")

# Faixas na MESMA ordem da Lista de Opções já cadastrada para o item
# (ListaOpcoes TECIDO_17073 / TECIDO_17074 já contêm todas as 14)
FAIXAS_ORDER = ['FX FORNECIDO','FX A','FX B','FX C','FX D','FX E','FX F',
                'FX G','FX H','FX I','FX J/N','FX R','FX V','FX KNIT']

# Preços da tabela mãe (PIU - DESIGNERS / MESC)
# 17073 = MESC SB AV 0,40 - 2,40X1,00 2ENC (linha 302)
precos_17073 = {
    'FX FORNECIDO': 5525, 'FX A': 5706, 'FX B': 5767, 'FX C': 5827,
    'FX D': 5912, 'FX E': 5973, 'FX F': 6069, 'FX G': 6154,
    'FX H': 6251, 'FX I': 6372, 'FX J/N': 6674, 'FX R': 6795,
    'FX V': 7037, 'FX KNIT': 8187,
}
# 17074 = MESC CSB 1,50X1,66 2ENC (linha 292)
precos_17074 = {
    'FX FORNECIDO': 5388, 'FX A': 5588, 'FX B': 5655, 'FX C': 5722,
    'FX D': 5815, 'FX E': 5882, 'FX F': 5989, 'FX G': 6082,
    'FX H': 6189, 'FX I': 6322, 'FX J/N': 6656, 'FX R': 6789,
    'FX V': 7056, 'FX KNIT': 8323,
}

itens = {
    '17073': ('240X100X90',  precos_17073),
    '17074': ('150X166X90H', precos_17074),
}

# 1) Localizar e SOBRESCREVER a regra errada existente de cada item
#    (vira a primeira faixa = FX FORNECIDO), depois anexar as demais.
for cod, (dim, precos) in itens.items():
    # encontra a linha da regra errada
    linha_existente = None
    for r in range(2, wsRP.max_row + 1):
        if str(wsRP.cell(r, 1).value or '') == cod:
            linha_existente = r
            break
    if linha_existente is None:
        print(f"  ERRO: {cod} sem regra existente!")
        continue

    car_dim = f'DIMENSAO_{cod}'
    car_tec = f'TECIDO_{cod}'

    # primeira faixa reaproveita a linha existente
    primeira = FAIXAS_ORDER[0]
    cond = f'{car_dim}="{dim}" E {car_tec}="{primeira}"'
    set_regra(wsRP, linha_existente, cod, cond, precos[primeira])
    print(f"  {cod}: regra errada reescrita -> {primeira} => {precos[primeira]}")

    # demais faixas anexadas ao final
    r_append = _ult(wsRP, 7) + 1
    for faixa in FAIXAS_ORDER[1:]:
        cond = f'{car_dim}="{dim}" E {car_tec}="{faixa}"'
        set_regra(wsRP, r_append, cod, cond, precos[faixa])
        print(f"  {cod}: + {faixa} => {precos[faixa]}")
        r_append += 1

# 2) Atualizar refs das tabelas (só Regras de Preço mudou de tamanho)
TABELAS = [
    ("Lista de Opções",          "Tabela_ListaOpcoes",      5, "E"),
    ("Características",          "Tabela_Caracteristicas",  4, "D"),
    ("Itens",                    "Tabela_Itens",           14, "N"),
    ("Itens - Atributos",        "Tabela_ItensAtributos",   3, "C"),
    ("Itens - Regras de Preço",  "Tabela_ItensRegrasPreco", 7, "G"),
]
for sheet, tab, nc, L in TABELAS:
    ws = wb[sheet]
    new_ref = f"A1:{L}{_ult(ws, nc)}"
    ws.tables[tab].ref = new_ref
    print(f"  Table {tab}: ref -> {new_ref}")

wb.save(DST)
print(f"\n✓ Saved {DST}")
