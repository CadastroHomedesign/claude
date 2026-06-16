"""
Correção 7: CACAU (POL 22042 / PUFF 22043) segue a tabela ITALO TROPICAL.
Prova: FX R e FX V do Focco são idênticos ao Italo. Só FORNECIDO e N divergiam.
Trabalha sobre PIU_MOBILE_HDA_CORRIGIDO6.xlsx -> PIU_MOBILE_HDA_CORRIGIDO7.xlsx
Altera apenas a coluna 6 (preço); B/D preservados.
"""
import openpyxl, shutil
SRC="PIU_MOBILE_HDA_CORRIGIDO6.xlsx"; DST="PIU_MOBILE_HDA_CORRIGIDO7.xlsx"
shutil.copy2(SRC,DST)
wb=openpyxl.load_workbook(DST,data_only=False); ws=wb["Itens - Regras de Preço"]

# (cod, grade-substring) -> valor Italo
fixes = {
    ('22042','TEC_FORNECIDO"'): 3804,   # FX FORN RB (aspas finais p/ não pegar _V_8001)
    ('22042','TECIDO_N"'):      4490,   # FX J/N
    ('22043','TECIDO_FORNECIDO"'): 2099,
    ('22043','TECIDO_N"'):      2351,
}
n=0
for r in range(2,ws.max_row+1):
    cod=str(ws.cell(r,1).value or ''); cond=str(ws.cell(r,3).value or '')
    for (c,sub),v in fixes.items():
        if cod==c and sub in cond:
            if ws.cell(r,6).value!=v:
                ws.cell(r,6).value=v; n+=1
                print(f"  {cod} {sub} -> {v}")
            break
print(f"\nTotal: {n}")

def _ult(w,nc):
    u=w.max_row
    while u>1 and all(w.cell(u,c).value is None for c in range(1,nc+1)): u-=1
    return u
for s,t,nc,L in [("Lista de Opções","Tabela_ListaOpcoes",5,"E"),
                 ("Características","Tabela_Caracteristicas",4,"D"),
                 ("Itens","Tabela_Itens",14,"N"),
                 ("Itens - Atributos","Tabela_ItensAtributos",3,"C"),
                 ("Itens - Regras de Preço","Tabela_ItensRegrasPreco",7,"G")]:
    w=wb[s]; w.tables[t].ref=f"A1:{L}{_ult(w,nc)}"
wb.save(DST); print(f"✓ Saved {DST}")
